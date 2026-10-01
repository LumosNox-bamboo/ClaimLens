"""Read verification outputs and join local names by exact candidate_id only."""

from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter, defaultdict
from copy import deepcopy
from pathlib import Path
from urllib.parse import unquote, urlsplit

from openpyxl import load_workbook

from claimlens.verification.models import Status, validate_candidate

PUBLIC_TYPES = ("publication", "award", "competition", "patent", "conference_presentation")
MATERIAL_TYPES = ("education", "certification", "professional_qualification")
REVIEW_STATUSES = ("CONFLICT", "PARTIALLY_VERIFIED", "NEEDS_REVIEW", "NOT_FOUND")
TYPE_LABELS = {
    "publication": "Publications",
    "award": "Awards",
    "competition": "Competitions",
    "patent": "Patents",
    "conference_presentation": "Conferences",
    "education": "Education",
    "certification": "Certifications",
    "professional_qualification": "Professional qualifications",
}
FACT_KEYS = (
    "title",
    "claim_text",
    "year",
    "journal",
    "doi",
    "authors",
    "organization",
    "patent_number",
    "award_name",
    "publication_status",
)
META_KEYS = (
    "title",
    "journal",
    "doi",
    "authors",
    "year",
    "volume",
    "issue",
    "page",
    "article_number",
    "publication_status",
    "patent_number",
)


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def warning(warnings, field, expected, observed, kind="DATA_INTEGRITY_WARNING"):
    difference = (
        observed - expected
        if isinstance(expected, (int, float)) and isinstance(observed, (int, float))
        else "values differ"
    )
    warnings.append(
        dict(type=kind, field=field, expected=expected, observed=observed, difference=difference)
    )


def read_mapping(path, warnings):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = next(rows)
        indices = {
            k: headers.index(k) for k in ("candidate_id", "application_id", "candidate_name")
        }
        mapping = {}
        names = set()
        ambiguous = set()
        for row in rows:
            cid = str(row[indices["candidate_id"]] or "").strip()
            if not cid:
                continue
            record = {k: str(row[i] or "").strip() for k, i in indices.items()}
            if record["candidate_name"]:
                names.add(record["candidate_name"])
            if cid in mapping:
                ambiguous.add(cid)
                warning(warnings, "duplicate_identity_mapping:" + cid, 1, 2)
                mapping[cid]["candidate_name"] = ""
            else:
                mapping[cid] = record
        for cid in ambiguous:
            mapping[cid]["candidate_name"] = ""
        return mapping, names
    finally:
        wb.close()


def clean_text(value):
    text = str(value)
    text = re.sub(
        r'(?:/Users/|/home/|/private/|[A-Za-z]:\\)[^\s<>"\']+', "[local path removed]", text
    )
    text = re.sub(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", "[contact removed]", text)
    text = re.sub(r"(?<!\w)(?:\+?\d[ -]*){11,15}(?!\w)", "[contact number removed]", text)
    text = re.sub(r"(?<!\w)\d{17}[\dXx](?!\w)", "[identity number removed]", text)
    text = re.sub(r"(?:地址|住址|Address)\s*[:：][^\n;；]+", "[address removed]", text, flags=re.I)
    return text


def safe_url(value):
    value = str(value or "").strip()
    try:
        p = urlsplit(value)
        if (
            p.scheme.lower() not in {"http", "https"}
            or not p.netloc
            or p.username
            or p.password
            or any(ord(c) < 32 for c in value)
        ):
            return None
    except ValueError:
        return None
    return value


def scrub(value, names=()):
    if isinstance(value, dict):
        return {
            k: scrub(v, names)
            for k, v in value.items()
            if k
            not in {
                "candidate_name",
                "source_folder",
                "source_file",
                "local_path",
                "verification_query",
            }
        }
    if isinstance(value, list):
        return [scrub(v, names) for v in value]
    if isinstance(value, str):
        text = clean_text(value)
        for name in sorted(names, key=len, reverse=True):
            text = re.sub(re.escape(name), "[local name removed]", text, flags=re.I)
        return text
    return value


def review_reason(r, audit):
    text = (
        r.get("verification_notes", "") + " " + str(audit.get("input_quality_issue", ""))
    ).lower()
    extraction = any(
        x in text
        for x in (
            "truncated",
            "heading",
            "non-factual",
            "non_factual",
            "incomplete citation",
            "parser artifact",
            "generic skill",
        )
    )
    if extraction:
        reason = (
            "Truncated or incomplete citation"
            if any(x in text for x in ("truncated", "incomplete citation"))
            else "Heading or non-factual extracted statement"
        )
        return "Extraction issue", reason, True
    if any(
        x in text
        for x in (
            "recipient",
            "presenter",
            "inventor",
            "authorship",
            "author ambiguity",
            "individual award",
            "individual result",
        )
    ):
        reason = (
            "Missing inventor linkage"
            if "inventor" in text or r["claim_type"] == "patent"
            else "Missing presenter linkage"
            if r["claim_type"] == "conference_presentation"
            else "Missing recipient or author linkage"
        )
        return "Identity linkage issue", reason, False
    if any(
        x in text
        for x in (
            "missing_patent_identifier",
            "no patent number",
            "incomplete metadata",
            "year ambiguity",
            "title mismatch",
            "journal ambiguity",
        )
    ):
        return "Metadata issue", "Missing or ambiguous core metadata", False
    if any(
        x in text
        for x in (
            "unavailable",
            "inaccessible",
            "ambiguous",
            "no reliable",
            "insufficient public",
            "provider",
            "did not yield",
        )
    ):
        return "Source issue", "Unavailable, ambiguous or insufficient public evidence", False
    return "Other", "Manual assessment required", False


def priority(r):
    status = r["verification_status"]
    if status == "CONFLICT":
        return 1
    if status == "PARTIALLY_VERIFIED":
        return 2
    if status == "NEEDS_REVIEW":
        return 3 if r["review_category"] == "Identity linkage issue" else 4
    return 5 if status == "NOT_FOUND" else 6


def display_metadata(r):
    original = r.get("original_claim", {})
    metadata = next(
        (
            e.get("verified_metadata", {})
            for e in r.get("evidence", [])
            if e.get("verified_metadata")
        ),
        {},
    )
    title = (
        metadata.get("title")
        or original.get("title")
        or original.get("award_name")
        or original.get("claim_text")
        or "No structured title supplied"
    )
    title = re.sub(r"^\s*\d+[.)、]\s*", "", str(title))
    if len(title) > 190:
        title = title[:187].rstrip() + "…"
    return dict(
        title=title,
        **{
            k: original.get(k) or metadata.get(k)
            for k in (
                "authors",
                "journal",
                "year",
                "doi",
                "publication_status",
                "organization",
                "patent_number",
            )
            if original.get(k) or metadata.get(k)
        },
    )


def report_claim(raw, audit):
    r = {
        k: deepcopy(raw.get(k))
        for k in (
            "claim_id",
            "claim_type",
            "verification_status",
            "matched_fields",
            "conflicting_fields",
            "unverified_fields",
            "evidence_strength",
            "freshness_status",
            "suggested_cv_update",
            "journal_metrics",
            "verification_notes",
        )
    }
    r["original_claim"] = {
        k: raw.get("original_claim", {}).get(k)
        for k in FACT_KEYS
        if raw.get("original_claim", {}).get(k)
    }
    r["evidence"] = [
        {
            k: deepcopy(e.get(k))
            for k in (
                "source_name",
                "source_type",
                "url",
                "accessed_at",
                "evidence_strength",
                "matched_fields",
                "conflicting_fields",
                "notes",
            )
        }
        | {
            "verified_metadata": {
                k: v for k, v in e.get("verified_metadata", {}).items() if k in META_KEYS
            }
        }
        for e in raw.get("evidence", [])
    ]
    r["review_category"], r["review_reason"], r["extraction_issue"] = review_reason(r, audit)
    r["review_priority"] = priority(r)
    r["display"] = display_metadata(r)
    r["audit_summary"] = dict(
        search_attempt_count=audit.get("search_attempt_count", 0),
        sources_checked=sorted(
            {
                urlsplit(s["url"]).netloc
                for s in audit.get("sources_considered", [])
                if safe_url(s.get("url"))
            }
        ),
    )
    return scrub(r)


def candidate(cid, application_id, name, claims, missing=False):
    statuses = Counter(c["verification_status"] for c in claims)
    types = Counter(c["claim_type"] for c in claims)
    return dict(
        candidate_id=cid,
        application_id=application_id,
        candidate_name=name,
        claims=claims,
        claims_total=len(claims),
        public_claims=sum(types[t] for t in PUBLIC_TYPES),
        material_claims=sum(types[t] for t in MATERIAL_TYPES),
        status_counts={s.value: statuses[s.value] for s in Status},
        claim_type_counts={t: types[t] for t in (*PUBLIC_TYPES, *MATERIAL_TYPES)},
        public_evidence_claims=sum(
            c["verification_status"] in ("VERIFIED", "PARTIALLY_VERIFIED") and bool(c["evidence"])
            for c in claims
        ),
        cv_updates=sum(c["freshness_status"] == "CV_UPDATE_AVAILABLE" for c in claims),
        high_value_review_count=sum(
            c["claim_type"] in PUBLIC_TYPES and c["verification_status"] in REVIEW_STATUSES
            for c in claims
        ),
        verification_result_missing=missing,
    )


def normalize_journal(name):
    return re.sub(r"[^\w]", "", unicodedata.normalize("NFKC", name).casefold())


def load_batch(batch, include_application_id=True):
    batch = Path(batch)
    source = batch / "02_VERIFICATION_RESULTS"
    expected = load_json(source / "batch_summary.json")
    warnings = []
    mapping, names = read_mapping(batch / "00_LOCAL_ONLY" / "candidate_map.xlsx", warnings)
    files = sorted(source.glob("C-*.verification.json"))
    candidates = []
    found = set()
    missing_mapping = []
    for path in files:
        cid = path.name.removesuffix(".verification.json")
        if not re.fullmatch(r"[A-Za-z0-9_-]+", cid):
            raise ValueError("Unsafe candidate identifier")
        d = load_json(path)
        if d.get("candidate_id") != cid:
            warning(warnings, "candidate_id:" + cid, cid, d.get("candidate_id"))
        try:
            validate_candidate(d)
        except (ValueError, TypeError, KeyError):
            warning(warnings, "verification_schema:" + cid, "valid", "invalid")
        found.add(cid)
        ap = source / "_logs" / (cid + ".audit.json")
        if ap.exists():
            a = load_json(ap)
            audits = {x["claim_id"]: x for x in a.get("claims", [])}
        else:
            audits = {}
            warning(warnings, "audit_log:" + cid, "present", "missing")
        identity = mapping.get(cid, {})
        if not identity:
            missing_mapping.append(cid)
            warning(warnings, "identity_mapping:" + cid, "present", "missing")
        app_id = str(d.get("application_id") or identity.get("application_id") or "")
        if identity.get("application_id") and d.get("application_id") != identity["application_id"]:
            warning(
                warnings,
                "application_id_match:" + cid,
                "matching local and verification identifiers",
                "mismatch",
            )
        claims = [report_claim(r, audits.get(r["claim_id"], {})) for r in d.get("results", [])]
        candidates.append(
            candidate(
                cid,
                app_id if include_application_id else "",
                identity.get("candidate_name", ""),
                claims,
            )
        )
    missing_results = sorted(set(mapping) - found)
    for cid in missing_results:
        if not re.fullmatch(r"[A-Za-z0-9_-]+", cid):
            raise ValueError("Unsafe mapping identifier")
        warning(warnings, "verification_result:" + cid, "present", "missing")
        identity = mapping[cid]
        candidates.append(
            candidate(
                cid,
                identity["application_id"] if include_application_id else "",
                identity["candidate_name"],
                [],
                True,
            )
        )
    all_claims = [r for c in candidates for r in c["claims"]]
    statuses = Counter(r["verification_status"] for r in all_claims)
    types = Counter(r["claim_type"] for r in all_claims)
    publications = [r for r in all_claims if r["claim_type"] == "publication"]
    jif = sum(
        bool(r["journal_metrics"] and r["journal_metrics"].get("latest_jif") is not None)
        for r in publications
    )
    failures = (
        len(load_json(source / "failed_items.json"))
        if (source / "failed_items.json").exists()
        else 0
    )
    observed = dict(
        total_candidates=len(files),
        completed_candidates=len(files),
        total_claims=len(all_claims),
        verification_status_counts={s.value: statuses[s.value] for s in Status},
        claim_type_counts=dict(types),
        CV_UPDATE_AVAILABLE=sum(c["cv_updates"] for c in candidates),
        JIF_found=jif,
        JIF_unavailable=len(publications) - jif,
        processing_failures=failures,
    )
    for key, value in observed.items():
        reference = expected.get(key)
        if isinstance(value, dict):
            for sub in set(value) | set(reference or {}):
                if value.get(sub, 0) != (reference or {}).get(sub, 0):
                    warning(
                        warnings, key + "." + sub, (reference or {}).get(sub, 0), value.get(sub, 0)
                    )
        elif reference is not None and reference != value:
            warning(warnings, key, reference, value)
    gaps = defaultdict(lambda: dict(publication_claims=0, candidates=set(), journal=""))
    unidentified = 0
    for c in candidates:
        for r in c["claims"]:
            if r["claim_type"] != "publication" or (
                r["journal_metrics"] and r["journal_metrics"].get("latest_jif") is not None
            ):
                continue
            journal = r["display"].get("journal")
            if not journal:
                unidentified += 1
                continue
            key = normalize_journal(journal)
            gaps[key]["journal"] = gaps[key]["journal"] or journal
            gaps[key]["publication_claims"] += 1
            gaps[key]["candidates"].add(c["candidate_id"])
    journal_gaps = dict(
        journals=[
            dict(
                journal=g["journal"],
                publication_claims=g["publication_claims"],
                candidate_count=len(g["candidates"]),
            )
            for _, g in sorted(gaps.items())
        ],
        unidentified_journal_claims=unidentified,
        grouping_rule="NFKC, casefold, punctuation/whitespace removed; unidentified journals counted separately, never guessed.",
    )
    dashboard = dict(
        **observed,
        candidates_rendered=len(candidates),
        public_claims=sum(c["public_claims"] for c in candidates),
        material_claims=sum(c["material_claims"] for c in candidates),
        high_value_review_queue_count=sum(c["high_value_review_count"] for c in candidates),
        extraction_issue_count=sum(r["extraction_issue"] for r in all_claims),
        data_integrity_warnings=warnings,
    )
    anonymous = scrub(candidates, names)
    # URLs must not expose mapped names through percent encoding.
    for c in anonymous:
        for r in c["claims"]:
            for e in r["evidence"]:
                if any(name.casefold() in unquote(e.get("url") or "").casefold() for name in names):
                    e["url"] = None
    return (
        candidates,
        anonymous,
        dashboard,
        journal_gaps,
        dict(identity_mapping_missing=missing_mapping, verification_result_missing=missing_results),
        names,
    )
