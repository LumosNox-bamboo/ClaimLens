from __future__ import annotations

from pathlib import Path
from typing import Protocol

from docx import Document
from pypdf import PdfReader


class DocumentParser(Protocol):
    def parse(self, path: Path) -> str: ...


class TextParser:
    def parse(self, path: Path) -> str:
        return path.read_text(encoding="utf-8", errors="replace")


class PDFTextParser:
    """Extract embedded text only. Scanned-PDF OCR is intentionally not automatic in v0.1."""

    def parse(self, path: Path) -> str:
        reader = PdfReader(str(path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if not text.strip():
            raise ValueError("No embedded PDF text found; OCR is not enabled in ClaimLens v0.1")
        return text


class DOCXParser:
    def parse(self, path: Path) -> str:
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(parts)


PARSERS: dict[str, DocumentParser] = {
    ".txt": TextParser(),
    ".pdf": PDFTextParser(),
    ".docx": DOCXParser(),
}


def extract_text(path: Path) -> str:
    parser = PARSERS.get(path.suffix.lower())
    if parser is None:
        raise ValueError(f"Unsupported file type: {path.suffix.lower()}")
    return parser.parse(path)


def embedded_image_count(path: Path) -> int:
    """Best-effort local image detection; never extracts or transmits image bytes."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from zipfile import ZipFile
        with ZipFile(path) as archive:
            return sum(name.startswith("word/media/") for name in archive.namelist())
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        count = 0
        for page in reader.pages:
            try:
                count += len(page.images)
            except Exception:
                # Image enumeration varies across PDFs; failure must not imply no privacy risk.
                return -1
        return count
    return 0
