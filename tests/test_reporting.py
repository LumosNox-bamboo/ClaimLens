import json

from openpyxl import Workbook

from claimlens.reporting.data import load_batch, safe_url
from claimlens.reporting.generate import generate, workbook_tables
from claimlens.reporting.render import candidate_page, claim_card
from claimlens.verification.aggregator import summary
from claimlens.verification.checkpoint import atomic_json
from claimlens.verification.evidence import evidence
from claimlens.verification.models import Status, result, now


def fixture_batch(tmp_path):
    source = tmp_path / "batch"
    (source / "00_LOCAL_ONLY").mkdir(parents=True)
    (source / "02_VERIFICATION_RESULTS/_logs").mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.append(["candidate_id", "application_id", "candidate_name", "source_folder"])
    ws.append(["C-SYNTH01", "SYNTH-001", "示例甲", "/synthetic/private/folder"])
    wb.save(source / "00_LOCAL_ONLY/candidate_map.xlsx")
    rs = []
    for i, status in enumerate(Status):
        c = dict(
            claim_id=f"SYNTH-CLAIM-{i}",
            claim_type="education" if status == Status.DOCUMENT_REQUIRED else "publication",
            title="A fictional title for synthetic review",
            year="2025",
            claim_text="A fictional title for synthetic review",
        )
        r = result(
            c,
            status,
            "Heading or non-factual statement"
            if status == Status.NEEDS_REVIEW
            else "Synthetic fixture",
        )
        if status in (Status.VERIFIED, Status.PARTIALLY_VERIFIED, Status.CONFLICT):
            r["matched_fields"] = ["title"]
            r["evidence"] = [
                evidence(
                    "Fictional publisher",
                    "official",
                    "https://example.org/work",
                    "A",
                    ["title"],
                    ["year"] if status == Status.CONFLICT else [],
                )
            ]
            r["evidence_strength"] = "A"
        if status == Status.CONFLICT:
            r["conflicting_fields"] = ["year"]
        if status == Status.VERIFIED:
            r["freshness_status"] = "CV_UPDATE_AVAILABLE"
            r["suggested_cv_update"] = {"publication_status": "published", "year": "2025"}
            r["journal_metrics"].update(
                latest_jif=4.8,
                jif_year=2025,
                metric_source="Fictional publisher",
                metric_source_url="https://example.org/metrics",
                metric_source_quality="official_publisher",
            )
        rs.append(r)
    d = dict(
        candidate_id="C-SYNTH01",
        application_id="SYNTH-001",
        verification_version="1.0",
        verified_at=now(),
        claims_total=len(rs),
        results=rs,
    )
    atomic_json(source / "02_VERIFICATION_RESULTS/C-SYNTH01.verification.json", d)
    atomic_json(
        source / "02_VERIFICATION_RESULTS/batch_summary.json",
        {**summary([d], 1, []), "run_complete": True},
    )
    atomic_json(source / "02_VERIFICATION_RESULTS/failed_items.json", [])
    atomic_json(
        source / "02_VERIFICATION_RESULTS/_logs/C-SYNTH01.audit.json",
        dict(
            claims=[
                dict(
                    claim_id=r["claim_id"],
                    search_attempt_count=2,
                    sources_considered=[{"url": "https://example.org/search"}],
                )
                for r in rs
            ]
        ),
    )
    return source


def test_named_and_anonymous_identity_boundary(tmp_path):
    source = fixture_batch(tmp_path)
    report = generate(source)
    out = source / "03_REPORTS"
    assert "示例甲" in (out / "named/C-SYNTH01.html").read_text()
    assert "示例甲" not in (out / "anonymous/C-SYNTH01.html").read_text()
    assert "示例甲" not in (out / "data/candidates_anonymous.json").read_text()
    assert "candidate_name" not in (out / "data/candidates_anonymous.json").read_text()
    assert "/synthetic/private/folder" not in (out / "anonymous/C-SYNTH01.html").read_text()
    assert report["source_files_unchanged"]
    assert not report["data_integrity_warnings"]


def test_materials_collapsed_and_absence_explanation(tmp_path):
    source = fixture_batch(tmp_path)
    generate(source)
    html = (source / "03_REPORTS/named/C-SYNTH01.html").read_text()
    assert '<details class="material-section">' in html
    assert '<details class="material-section" open' not in html
    assert "This does not indicate that the claim is false." in html
    assert "Public-source verification is not appropriate or sufficient" in html


def test_conflict_update_and_metrics_rendering(tmp_path):
    source = fixture_batch(tmp_path)
    named, *_ = load_batch(source)
    html = candidate_page(named[0])
    assert (
        "conflict-banner" in html and "Public evidence conflicts with the submitted claim" in html
    )
    assert "CV update available" in html and "Suggested update" in html
    assert "JIF 4.8 · 2025" in html
    missing = next(r for r in named[0]["claims"] if r["verification_status"] == "NOT_FOUND")
    card = claim_card(missing)
    assert "badge metric" not in card and "Journal metric not independently retrieved." in card
    assert 'target="_blank" rel="noopener noreferrer"' in html


def test_html_injection_and_unsafe_urls(tmp_path):
    source = fixture_batch(tmp_path)
    named, *_ = load_batch(source)
    r = named[0]["claims"][0]
    r["display"]["title"] = "<img src=x onerror=alert(1)>"
    r["evidence"][0]["url"] = "javascript:alert(1)"
    html = claim_card(r)
    assert "<img src=x" not in html and "&lt;img src=x" in html
    assert "javascript:" not in html
    assert safe_url("data:text/html,test") is None
    assert safe_url("https://user:secret@example.org/") is None


def test_mapping_and_missing_result_warnings(tmp_path):
    source = fixture_batch(tmp_path)
    wb = Workbook()
    ws = wb.active
    ws.append(["candidate_id", "application_id", "candidate_name"])
    ws.append(["C-SYNTH02", "SYNTH-002", "Example Candidate"])
    wb.save(source / "00_LOCAL_ONLY/candidate_map.xlsx")
    named, anonymous, dashboard, gaps, missing, names = load_batch(source)
    assert missing["identity_mapping_missing"] == ["C-SYNTH01"]
    assert missing["verification_result_missing"] == ["C-SYNTH02"]
    assert len(named) == 2 and any(c["verification_result_missing"] for c in named)
    assert all(w["type"] == "DATA_INTEGRITY_WARNING" for w in dashboard["data_integrity_warnings"])


def test_review_queue_excludes_document_claims_and_no_names_in_anonymous_tables(tmp_path):
    source = fixture_batch(tmp_path)
    named, anonymous, *_ = load_batch(source)
    tables = workbook_tables(named)
    headers = tables["Review Queue"]["headers"]
    index = headers.index("verification_status")
    assert all(row[index] != "DOCUMENT_REQUIRED" for row in tables["Review Queue"]["rows"])
    assert len(tables["Material Claims"]["rows"]) == 1
    anon_tables = workbook_tables(anonymous, True)
    assert "示例甲" not in json.dumps(anon_tables, ensure_ascii=False)
    assert all("candidate_name" not in t["headers"] for t in anon_tables.values())


def test_count_mismatch_is_visible_with_difference(tmp_path):
    source = fixture_batch(tmp_path)
    p = source / "02_VERIFICATION_RESULTS/batch_summary.json"
    d = json.loads(p.read_text())
    d["total_claims"] = 999
    atomic_json(p, d)
    report = generate(source)
    warning = next(w for w in report["data_integrity_warnings"] if w["field"] == "total_claims")
    assert warning["expected"] == 999 and warning["observed"] == 6 and warning["difference"] == -993
    assert "DATA_INTEGRITY_WARNING" in (source / "03_REPORTS/index.html").read_text()


def test_names_embedded_in_claims_are_scrubbed(tmp_path):
    source = fixture_batch(tmp_path)
    p = source / "02_VERIFICATION_RESULTS/C-SYNTH01.verification.json"
    d = json.loads(p.read_text())
    d["results"][0]["original_claim"]["claim_text"] = "示例甲 authored a fictional synthetic claim"
    atomic_json(p, d)
    generate(source)
    assert "示例甲" not in (source / "03_REPORTS/anonymous/C-SYNTH01.html").read_text()


def test_application_id_can_be_disabled(tmp_path):
    source = fixture_batch(tmp_path)
    generate(source, include_application_id=False)
    assert "SYNTH-001" not in (source / "03_REPORTS/anonymous/C-SYNTH01.html").read_text()


def test_duplicate_identity_mapping_does_not_choose_a_name(tmp_path):
    source = fixture_batch(tmp_path)
    from openpyxl import load_workbook

    path = source / "00_LOCAL_ONLY/candidate_map.xlsx"
    wb = load_workbook(path)
    wb.active.append(["C-SYNTH01", "SYNTH-002", "Example Candidate", ""])
    wb.save(path)
    wb.close()
    named, anonymous, dashboard, *_ = load_batch(source)
    assert named[0]["candidate_name"] == ""
    assert any(
        w["field"].startswith("duplicate_identity_mapping:")
        for w in dashboard["data_integrity_warnings"]
    )
    assert "Example Candidate" not in json.dumps(anonymous)


def test_journal_metric_gaps_deduplicate_normalized_journals(tmp_path):
    source = fixture_batch(tmp_path)
    path = source / "02_VERIFICATION_RESULTS/C-SYNTH01.verification.json"
    data = json.loads(path.read_text())
    for r in data["results"]:
        if r["claim_type"] == "publication":
            r["original_claim"]["journal"] = (
                "Fictional Journal"
                if r["verification_status"] == "NOT_FOUND"
                else "FICTIONAL JOURNAL"
            )
    atomic_json(path, data)
    _, _, _, gaps, *_ = load_batch(source)
    assert len(gaps["journals"]) == 1
    assert gaps["journals"][0]["publication_claims"] == 4
    assert gaps["journals"][0]["candidate_count"] == 1
